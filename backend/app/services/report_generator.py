"""
Report Generator Service.

Generates professional reports in multiple formats:
- PDF with charts and branding
- Excel/XLSX with formatted tables
- CSV for data export

Supports various report types:
- Daily/Weekly/Monthly summaries
- Doctor performance metrics
- Revenue reports
- Patient demographics
- No-show analysis
"""

import csv
import io
import logging
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, Optional
from uuid import UUID

import matplotlib
matplotlib.use('Agg')  # Use non-interactive backend
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
    PageBreak,
    Image,
)
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.clinic import Clinic
from app.models.scheduled_report import ReportType, ReportFormat
from app.services.analytics import AnalyticsService

logger = logging.getLogger(__name__)


@dataclass
class ReportData:
    """Container for report data."""
    title: str
    subtitle: Optional[str]
    clinic_name: str
    date_range: str
    sections: list[dict[str, Any]]
    metadata: dict[str, Any]


class ReportGenerator:
    """
    Report generation service.

    Generates professional reports with charts and branding.
    """

    def __init__(self, db: AsyncSession):
        """Initialize report generator with database session."""
        self.db = db
        self.analytics = AnalyticsService(db)

    async def generate_report(
        self,
        clinic_id: UUID,
        report_type: str,
        start_date: date,
        end_date: date,
        format: str = "pdf",
        parameters: Optional[dict] = None,
    ) -> bytes:
        """
        Generate a report in the specified format.

        Args:
            clinic_id: Clinic to generate report for
            report_type: Type of report (daily_summary, revenue_report, etc.)
            start_date: Report period start
            end_date: Report period end
            format: Output format (pdf, xlsx, csv)
            parameters: Additional parameters (doctor_id, etc.)

        Returns:
            Report file content as bytes
        """
        parameters = parameters or {}

        # Get clinic info
        clinic = await self.db.get(Clinic, clinic_id)
        if not clinic:
            raise ValueError(f"Clinic {clinic_id} not found")

        # Collect report data
        report_data = await self._collect_report_data(
            clinic=clinic,
            report_type=report_type,
            start_date=start_date,
            end_date=end_date,
            parameters=parameters,
        )

        # Generate in requested format
        if format == ReportFormat.PDF.value:
            return await self._generate_pdf(report_data)
        elif format == ReportFormat.EXCEL.value:
            return await self._generate_excel(report_data)
        elif format == ReportFormat.CSV.value:
            return await self._generate_csv(report_data)
        else:
            raise ValueError(f"Unsupported format: {format}")

    async def _collect_report_data(
        self,
        clinic: Clinic,
        report_type: str,
        start_date: date,
        end_date: date,
        parameters: dict,
    ) -> ReportData:
        """Collect data for the specified report type."""
        date_range = f"{start_date.strftime('%d %b %Y')} - {end_date.strftime('%d %b %Y')}"

        if report_type == ReportType.DAILY_SUMMARY.value:
            return await self._collect_daily_summary(clinic, start_date, end_date)
        elif report_type == ReportType.WEEKLY_SUMMARY.value:
            return await self._collect_weekly_summary(clinic, start_date, end_date)
        elif report_type == ReportType.MONTHLY_SUMMARY.value:
            return await self._collect_monthly_summary(clinic, start_date, end_date)
        elif report_type == ReportType.DOCTOR_PERFORMANCE.value:
            return await self._collect_doctor_performance(clinic, start_date, end_date, parameters)
        elif report_type == ReportType.REVENUE_REPORT.value:
            return await self._collect_revenue_report(clinic, start_date, end_date)
        elif report_type == ReportType.PATIENT_DEMOGRAPHICS.value:
            return await self._collect_patient_demographics(clinic, start_date, end_date)
        elif report_type == ReportType.NO_SHOW_ANALYSIS.value:
            return await self._collect_no_show_analysis(clinic, start_date, end_date)
        else:
            raise ValueError(f"Unknown report type: {report_type}")

    async def _collect_daily_summary(
        self, clinic: Clinic, start_date: date, end_date: date
    ) -> ReportData:
        """Collect data for daily summary report."""
        appt_stats = await self.analytics.get_appointment_stats(
            clinic.id, start_date, end_date
        )
        revenue_stats = await self.analytics.get_revenue_stats(
            clinic.id, start_date, end_date
        )

        return ReportData(
            title="Daily Practice Summary",
            subtitle=f"Report for {start_date.strftime('%d %B %Y')}",
            clinic_name=clinic.name,
            date_range=start_date.strftime('%d %b %Y'),
            sections=[
                {
                    "title": "Appointment Statistics",
                    "type": "table",
                    "headers": ["Metric", "Count", "Percentage"],
                    "rows": [
                        ["Total Appointments", appt_stats.total, "100%"],
                        ["Completed", appt_stats.completed, f"{appt_stats.completion_rate:.1f}%"],
                        ["Scheduled", appt_stats.scheduled, f"{(appt_stats.scheduled/appt_stats.total*100) if appt_stats.total > 0 else 0:.1f}%"],
                        ["Cancelled", appt_stats.cancelled, f"{appt_stats.cancellation_rate:.1f}%"],
                        ["No-Show", appt_stats.no_show, f"{appt_stats.no_show_rate:.1f}%"],
                    ],
                },
                {
                    "title": "Revenue Summary",
                    "type": "table",
                    "headers": ["Metric", "Amount (₹)"],
                    "rows": [
                        ["Total Revenue", f"{revenue_stats.total_revenue:,.2f}"],
                        ["Collected", f"{revenue_stats.collected:,.2f}"],
                        ["Pending", f"{revenue_stats.pending:,.2f}"],
                        ["Collection Rate", f"{revenue_stats.collection_rate:.1f}%"],
                        ["Average Invoice", f"{revenue_stats.average_invoice:,.2f}"],
                    ],
                },
            ],
            metadata={
                "generated_at": datetime.now().isoformat(),
                "period": "daily",
            },
        )

    async def _collect_weekly_summary(
        self, clinic: Clinic, start_date: date, end_date: date
    ) -> ReportData:
        """Collect data for weekly summary report."""
        appt_stats = await self.analytics.get_appointment_stats(
            clinic.id, start_date, end_date
        )
        revenue_stats = await self.analytics.get_revenue_stats(
            clinic.id, start_date, end_date
        )
        daily_metrics = await self.analytics.get_daily_metrics(
            clinic.id, start_date, end_date
        )
        doctor_utils = await self.analytics.get_doctor_utilization(
            clinic.id, start_date, end_date
        )

        return ReportData(
            title="Weekly Practice Summary",
            subtitle=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            clinic_name=clinic.name,
            date_range=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            sections=[
                {
                    "title": "Appointment Statistics",
                    "type": "table",
                    "headers": ["Metric", "Count", "Percentage"],
                    "rows": [
                        ["Total Appointments", appt_stats.total, "100%"],
                        ["Completed", appt_stats.completed, f"{appt_stats.completion_rate:.1f}%"],
                        ["Scheduled", appt_stats.scheduled, f"{(appt_stats.scheduled/appt_stats.total*100) if appt_stats.total > 0 else 0:.1f}%"],
                        ["Cancelled", appt_stats.cancelled, f"{appt_stats.cancellation_rate:.1f}%"],
                        ["No-Show", appt_stats.no_show, f"{appt_stats.no_show_rate:.1f}%"],
                    ],
                },
                {
                    "title": "Revenue Summary",
                    "type": "table",
                    "headers": ["Metric", "Amount (₹)"],
                    "rows": [
                        ["Total Revenue", f"{revenue_stats.total_revenue:,.2f}"],
                        ["Collected", f"{revenue_stats.collected:,.2f}"],
                        ["Pending", f"{revenue_stats.pending:,.2f}"],
                        ["Collection Rate", f"{revenue_stats.collection_rate:.1f}%"],
                    ],
                },
                {
                    "title": "Daily Trend",
                    "type": "chart",
                    "chart_type": "line",
                    "data": {
                        "labels": [m.date.strftime('%a %d') for m in daily_metrics],
                        "datasets": [
                            {
                                "label": "Appointments",
                                "data": [m.appointments for m in daily_metrics],
                            },
                        ],
                    },
                },
                {
                    "title": "Doctor Utilization",
                    "type": "table",
                    "headers": ["Doctor", "Total Slots", "Booked", "Utilization %"],
                    "rows": [
                        [
                            u.doctor_name,
                            u.total_slots,
                            u.booked_slots,
                            f"{u.utilization_rate:.1f}%",
                        ]
                        for u in doctor_utils[:10]  # Top 10
                    ],
                },
            ],
            metadata={
                "generated_at": datetime.now().isoformat(),
                "period": "weekly",
            },
        )

    async def _collect_monthly_summary(
        self, clinic: Clinic, start_date: date, end_date: date
    ) -> ReportData:
        """Collect data for monthly summary report."""
        appt_stats = await self.analytics.get_appointment_stats(
            clinic.id, start_date, end_date
        )
        revenue_stats = await self.analytics.get_revenue_stats(
            clinic.id, start_date, end_date
        )
        daily_metrics = await self.analytics.get_daily_metrics(
            clinic.id, start_date, end_date
        )
        doctor_utils = await self.analytics.get_doctor_utilization(
            clinic.id, start_date, end_date
        )
        demographics = await self.analytics.get_patient_demographics(
            clinic.id, start_date
        )

        return ReportData(
            title="Monthly Practice Summary",
            subtitle=f"{start_date.strftime('%B %Y')}",
            clinic_name=clinic.name,
            date_range=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            sections=[
                {
                    "title": "Key Performance Indicators",
                    "type": "table",
                    "headers": ["Metric", "Value"],
                    "rows": [
                        ["Total Appointments", appt_stats.total],
                        ["Completion Rate", f"{appt_stats.completion_rate:.1f}%"],
                        ["Total Revenue", f"₹{revenue_stats.total_revenue:,.2f}"],
                        ["Collection Rate", f"{revenue_stats.collection_rate:.1f}%"],
                        ["New Patients", demographics.new_this_period],
                        ["Average Invoice", f"₹{revenue_stats.average_invoice:,.2f}"],
                    ],
                },
                {
                    "title": "Revenue Breakdown",
                    "type": "table",
                    "headers": ["Category", "Amount (₹)"],
                    "rows": [
                        ["Total Billed", f"{revenue_stats.total_revenue:,.2f}"],
                        ["Collected", f"{revenue_stats.collected:,.2f}"],
                        ["Pending", f"{revenue_stats.pending:,.2f}"],
                    ],
                },
                {
                    "title": "Appointment Trend",
                    "type": "chart",
                    "chart_type": "line",
                    "data": {
                        "labels": [m.date.strftime('%d %b') for m in daily_metrics],
                        "datasets": [
                            {
                                "label": "Daily Appointments",
                                "data": [m.appointments for m in daily_metrics],
                            },
                        ],
                    },
                },
                {
                    "title": "Top Performing Doctors",
                    "type": "table",
                    "headers": ["Doctor", "Appointments", "Revenue (₹)", "Utilization %"],
                    "rows": [
                        [
                            u.doctor_name,
                            u.completed_appointments,
                            f"{u.revenue_generated:,.2f}",
                            f"{u.utilization_rate:.1f}%",
                        ]
                        for u in sorted(doctor_utils, key=lambda x: x.completed_appointments, reverse=True)[:10]
                    ],
                },
            ],
            metadata={
                "generated_at": datetime.now().isoformat(),
                "period": "monthly",
            },
        )

    async def _collect_doctor_performance(
        self, clinic: Clinic, start_date: date, end_date: date, parameters: dict
    ) -> ReportData:
        """Collect data for doctor performance report."""
        doctor_id = parameters.get("doctor_id")
        doctor_utils = await self.analytics.get_doctor_utilization(
            clinic.id, start_date, end_date
        )

        if doctor_id:
            doctor_utils = [u for u in doctor_utils if str(u.doctor_id) == str(doctor_id)]

        return ReportData(
            title="Doctor Performance Report",
            subtitle=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            clinic_name=clinic.name,
            date_range=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            sections=[
                {
                    "title": "Doctor Metrics",
                    "type": "table",
                    "headers": ["Doctor", "Total Slots", "Booked", "Completed", "Utilization %", "Avg Duration (min)", "Revenue (₹)"],
                    "rows": [
                        [
                            u.doctor_name,
                            u.total_slots,
                            u.booked_slots,
                            u.completed_appointments,
                            f"{u.utilization_rate:.1f}%",
                            f"{u.average_duration_minutes:.1f}",
                            f"{u.revenue_generated:,.2f}",
                        ]
                        for u in doctor_utils
                    ],
                },
            ],
            metadata={
                "generated_at": datetime.now().isoformat(),
                "period": "custom",
                "doctor_id": doctor_id,
            },
        )

    async def _collect_revenue_report(
        self, clinic: Clinic, start_date: date, end_date: date
    ) -> ReportData:
        """Collect data for revenue report."""
        revenue_stats = await self.analytics.get_revenue_stats(
            clinic.id, start_date, end_date
        )
        daily_metrics = await self.analytics.get_daily_metrics(
            clinic.id, start_date, end_date
        )

        return ReportData(
            title="Revenue Report",
            subtitle=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            clinic_name=clinic.name,
            date_range=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            sections=[
                {
                    "title": "Revenue Summary",
                    "type": "table",
                    "headers": ["Metric", "Amount (₹)"],
                    "rows": [
                        ["Total Revenue", f"{revenue_stats.total_revenue:,.2f}"],
                        ["Collected", f"{revenue_stats.collected:,.2f}"],
                        ["Pending", f"{revenue_stats.pending:,.2f}"],
                        ["Collection Rate", f"{revenue_stats.collection_rate:.1f}%"],
                        ["Average Invoice", f"{revenue_stats.average_invoice:,.2f}"],
                    ],
                },
                {
                    "title": "Daily Revenue Trend",
                    "type": "chart",
                    "chart_type": "bar",
                    "data": {
                        "labels": [m.date.strftime('%d %b') for m in daily_metrics],
                        "datasets": [
                            {
                                "label": "Revenue (₹)",
                                "data": [float(m.revenue) for m in daily_metrics],
                            },
                        ],
                    },
                },
            ],
            metadata={
                "generated_at": datetime.now().isoformat(),
                "period": "custom",
            },
        )

    async def _collect_patient_demographics(
        self, clinic: Clinic, start_date: date, end_date: date
    ) -> ReportData:
        """Collect data for patient demographics report."""
        demographics = await self.analytics.get_patient_demographics(
            clinic.id, start_date
        )

        return ReportData(
            title="Patient Demographics Report",
            subtitle=f"As of {end_date.strftime('%d %B %Y')}",
            clinic_name=clinic.name,
            date_range=f"New patients since {start_date.strftime('%d %b %Y')}",
            sections=[
                {
                    "title": "Patient Summary",
                    "type": "table",
                    "headers": ["Metric", "Count"],
                    "rows": [
                        ["Total Patients", demographics.total_patients],
                        ["New This Period", demographics.new_this_period],
                    ],
                },
                {
                    "title": "Gender Distribution",
                    "type": "table",
                    "headers": ["Gender", "Count", "Percentage"],
                    "rows": [
                        [
                            gender.title(),
                            count,
                            f"{(count/demographics.total_patients*100) if demographics.total_patients > 0 else 0:.1f}%",
                        ]
                        for gender, count in demographics.gender_breakdown.items()
                    ],
                },
                {
                    "title": "City Distribution",
                    "type": "table",
                    "headers": ["City", "Count"],
                    "rows": [
                        [city.title(), count]
                        for city, count in sorted(
                            demographics.city_breakdown.items(),
                            key=lambda x: x[1],
                            reverse=True,
                        )[:10]  # Top 10 cities
                    ],
                },
            ],
            metadata={
                "generated_at": datetime.now().isoformat(),
                "period": "all_time",
            },
        )

    async def _collect_no_show_analysis(
        self, clinic: Clinic, start_date: date, end_date: date
    ) -> ReportData:
        """Collect data for no-show analysis report."""
        analysis = await self.analytics.get_no_show_analysis(
            clinic.id, start_date, end_date
        )

        return ReportData(
            title="No-Show Analysis Report",
            subtitle=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            clinic_name=clinic.name,
            date_range=f"{start_date.strftime('%d %b')} - {end_date.strftime('%d %b %Y')}",
            sections=[
                {
                    "title": "No-Show Summary",
                    "type": "table",
                    "headers": ["Metric", "Value"],
                    "rows": [
                        ["Total No-Shows", analysis["total_no_shows"]],
                        ["Worst Day", analysis["worst_day"] or "N/A"],
                        ["Worst Hour", f"{analysis['worst_hour']}:00" if analysis["worst_hour"] else "N/A"],
                    ],
                },
                {
                    "title": "No-Shows by Day of Week",
                    "type": "table",
                    "headers": ["Day", "Count"],
                    "rows": [
                        [day, count]
                        for day, count in analysis["by_day_of_week"].items()
                    ],
                },
                {
                    "title": "No-Shows by Hour",
                    "type": "chart",
                    "chart_type": "bar",
                    "data": {
                        "labels": [f"{h}:00" for h in range(24)],
                        "datasets": [
                            {
                                "label": "No-Shows",
                                "data": [analysis["by_hour"].get(h, 0) for h in range(24)],
                            },
                        ],
                    },
                },
            ],
            metadata={
                "generated_at": datetime.now().isoformat(),
                "period": "custom",
            },
        )

    async def _generate_pdf(self, report_data: ReportData) -> bytes:
        """Generate PDF report using ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=72,
        )

        # Container for the PDF elements
        elements = []
        styles = getSampleStyleSheet()

        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#1a73e8'),
            spaceAfter=12,
            alignment=1,  # Center
        )
        subtitle_style = ParagraphStyle(
            'CustomSubtitle',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#5f6368'),
            spaceAfter=6,
            alignment=1,  # Center
        )
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#1a73e8'),
            spaceAfter=12,
            spaceBefore=20,
        )

        # Title
        elements.append(Paragraph(report_data.title, title_style))
        if report_data.subtitle:
            elements.append(Paragraph(report_data.subtitle, subtitle_style))
        elements.append(Paragraph(f"<b>{report_data.clinic_name}</b>", subtitle_style))
        elements.append(Spacer(1, 0.3 * inch))

        # Add sections
        for section in report_data.sections:
            # Section title
            elements.append(Paragraph(section["title"], heading_style))

            if section["type"] == "table":
                # Create table
                table_data = [section["headers"]] + section["rows"]
                table = Table(table_data, repeatRows=1)

                # Style the table
                table.setStyle(TableStyle([
                    # Header
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a73e8')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    # Body
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 9),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('GRID', (0, 0), (-1, -1), 1, colors.grey),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
                ]))

                elements.append(table)
                elements.append(Spacer(1, 0.3 * inch))

            elif section["type"] == "chart":
                # Generate chart using matplotlib
                chart_path = self._generate_chart(section)
                if chart_path and Path(chart_path).exists():
                    img = Image(chart_path, width=5*inch, height=3*inch)
                    elements.append(img)
                    elements.append(Spacer(1, 0.3 * inch))
                    # Clean up
                    Path(chart_path).unlink()

        # Footer
        elements.append(Spacer(1, 0.5 * inch))
        footer_style = ParagraphStyle(
            'Footer',
            parent=styles['Normal'],
            fontSize=8,
            textColor=colors.grey,
            alignment=1,  # Center
        )
        elements.append(Paragraph(
            f"Generated on {datetime.now().strftime('%d %B %Y at %H:%M')} | DocAssist Practice Manager",
            footer_style
        ))

        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        return buffer.read()

    def _generate_chart(self, section: dict) -> Optional[str]:
        """Generate a chart using matplotlib."""
        try:
            data = section["data"]
            chart_type = section.get("chart_type", "line")

            plt.figure(figsize=(10, 6))

            for dataset in data["datasets"]:
                if chart_type == "line":
                    plt.plot(data["labels"], dataset["data"], marker='o', label=dataset["label"])
                elif chart_type == "bar":
                    plt.bar(data["labels"], dataset["data"], label=dataset["label"])

            plt.xlabel("")
            plt.ylabel("")
            plt.title(section["title"])
            plt.legend()
            plt.xticks(rotation=45, ha='right')
            plt.tight_layout()

            # Save to temp file
            chart_path = f"/tmp/chart_{datetime.now().timestamp()}.png"
            plt.savefig(chart_path, dpi=150, bbox_inches='tight')
            plt.close()

            return chart_path
        except Exception as e:
            logger.error(f"Failed to generate chart: {e}")
            return None

    async def _generate_excel(self, report_data: ReportData) -> bytes:
        """Generate Excel report using openpyxl."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"

        # Title
        ws.merge_cells('A1:E1')
        title_cell = ws['A1']
        title_cell.value = report_data.title
        title_cell.font = Font(size=18, bold=True, color="1a73e8")
        title_cell.alignment = Alignment(horizontal='center')

        # Subtitle
        if report_data.subtitle:
            ws.merge_cells('A2:E2')
            subtitle_cell = ws['A2']
            subtitle_cell.value = report_data.subtitle
            subtitle_cell.font = Font(size=12, color="5f6368")
            subtitle_cell.alignment = Alignment(horizontal='center')

        # Clinic name
        ws.merge_cells('A3:E3')
        clinic_cell = ws['A3']
        clinic_cell.value = report_data.clinic_name
        clinic_cell.font = Font(size=12, bold=True)
        clinic_cell.alignment = Alignment(horizontal='center')

        current_row = 5

        # Add sections
        for section in report_data.sections:
            # Section title
            ws.merge_cells(f'A{current_row}:E{current_row}')
            section_title = ws[f'A{current_row}']
            section_title.value = section["title"]
            section_title.font = Font(size=14, bold=True, color="1a73e8")
            current_row += 1

            if section["type"] == "table":
                # Headers
                header_fill = PatternFill(start_color="1a73e8", end_color="1a73e8", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")

                for col_idx, header in enumerate(section["headers"], 1):
                    cell = ws.cell(row=current_row, column=col_idx)
                    cell.value = header
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal='center')

                current_row += 1

                # Data rows
                row_fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
                row_fill_gray = PatternFill(start_color="f5f5f5", end_color="f5f5f5", fill_type="solid")

                for row_idx, row_data in enumerate(section["rows"]):
                    fill = row_fill_white if row_idx % 2 == 0 else row_fill_gray
                    for col_idx, value in enumerate(row_data, 1):
                        cell = ws.cell(row=current_row, column=col_idx)
                        cell.value = value
                        cell.fill = fill
                        cell.alignment = Alignment(horizontal='left')
                    current_row += 1

            current_row += 2  # Space between sections

        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save to bytes
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.read()

    async def _generate_csv(self, report_data: ReportData) -> bytes:
        """Generate CSV report."""
        buffer = io.StringIO()
        writer = csv.writer(buffer)

        # Title and metadata
        writer.writerow([report_data.title])
        if report_data.subtitle:
            writer.writerow([report_data.subtitle])
        writer.writerow([report_data.clinic_name])
        writer.writerow([report_data.date_range])
        writer.writerow([])

        # Add sections
        for section in report_data.sections:
            if section["type"] == "table":
                writer.writerow([section["title"]])
                writer.writerow(section["headers"])
                writer.writerows(section["rows"])
                writer.writerow([])  # Empty row between sections

        buffer.seek(0)
        return buffer.getvalue().encode('utf-8')


def get_report_generator(db: AsyncSession) -> ReportGenerator:
    """Factory function for report generator."""
    return ReportGenerator(db)
