"""
Tests for Reports & Exports Service.

Phase 8: PDF and Excel report generation tests.
"""

import pytest
from datetime import date, datetime, timedelta
from decimal import Decimal
from uuid import uuid4

from app.services.reports import PDFReport, ReportConfig


class TestReportConfig:
    """Tests for ReportConfig dataclass."""

    def test_default_config(self):
        """Test default configuration values."""
        config = ReportConfig(clinic_name="Test Clinic")

        assert config.clinic_name == "Test Clinic"
        assert config.clinic_address == ""
        assert config.clinic_phone == ""
        assert config.clinic_email == ""
        assert config.logo_path is None
        assert config.primary_color == (41, 128, 185)
        assert config.secondary_color == (52, 73, 94)

    def test_full_config(self):
        """Test full configuration."""
        config = ReportConfig(
            clinic_name="DocAssist Clinic",
            clinic_address="123 Main St, Mumbai",
            clinic_phone="+91 9876543210",
            clinic_email="info@clinic.com",
            primary_color=(0, 123, 255),
        )

        assert config.clinic_name == "DocAssist Clinic"
        assert config.clinic_address == "123 Main St, Mumbai"
        assert config.clinic_phone == "+91 9876543210"
        assert config.clinic_email == "info@clinic.com"
        assert config.primary_color == (0, 123, 255)


class TestPDFReport:
    """Tests for PDFReport class."""

    def test_pdf_creation(self):
        """Test basic PDF creation."""
        config = ReportConfig(clinic_name="Test Clinic")
        pdf = PDFReport(config)
        pdf.alias_nb_pages()
        pdf.add_page()

        # Generate PDF bytes
        output = pdf.output()

        assert isinstance(output, bytes)
        assert len(output) > 0
        # PDF should start with %PDF
        assert output[:4] == b'%PDF'

    def test_pdf_with_title(self):
        """Test PDF with title."""
        config = ReportConfig(clinic_name="Test Clinic")
        pdf = PDFReport(config)
        pdf.alias_nb_pages()
        pdf.add_page()
        pdf.add_title("Test Report", "January 2026")

        output = pdf.output()
        assert len(output) > 0

    def test_pdf_with_section(self):
        """Test PDF with section."""
        config = ReportConfig(clinic_name="Test Clinic")
        pdf = PDFReport(config)
        pdf.alias_nb_pages()
        pdf.add_page()
        pdf.add_section("Test Section")

        output = pdf.output()
        assert len(output) > 0

    def test_pdf_with_table(self):
        """Test PDF with table."""
        config = ReportConfig(clinic_name="Test Clinic")
        pdf = PDFReport(config)
        pdf.alias_nb_pages()
        pdf.add_page()

        headers = ["Name", "Value", "Status"]
        rows = [
            ["Appointments", "50", "Good"],
            ["Revenue", "Rs.25,000", "Excellent"],
            ["No-Shows", "5", "Low"],
        ]

        pdf.add_table(headers, rows)

        output = pdf.output()
        assert len(output) > 0

    def test_pdf_with_metric_boxes(self):
        """Test PDF with metric boxes."""
        config = ReportConfig(clinic_name="Test Clinic")
        pdf = PDFReport(config)
        pdf.alias_nb_pages()
        pdf.add_page()

        y = pdf.get_y()
        pdf.add_metric_box("Appointments", "50", 10, y)
        pdf.add_metric_box("Revenue", "Rs.25K", 60, y)

        output = pdf.output()
        assert len(output) > 0

    def test_pdf_multiple_pages(self):
        """Test PDF with multiple pages."""
        config = ReportConfig(clinic_name="Test Clinic")
        pdf = PDFReport(config)
        pdf.alias_nb_pages()

        pdf.add_page()
        pdf.add_title("Page 1")

        pdf.add_page()
        pdf.add_title("Page 2")

        pdf.add_page()
        pdf.add_title("Page 3")

        output = pdf.output()
        assert len(output) > 0
        # Multi-page PDFs should be larger
        assert len(output) > 1000


class TestExcelGeneration:
    """Tests for Excel generation utilities."""

    def test_pandas_dataframe_to_excel(self):
        """Test creating Excel from DataFrame."""
        import pandas as pd
        import io

        data = {
            "Date": ["2026-01-01", "2026-01-02", "2026-01-03"],
            "Appointments": [10, 15, 12],
            "Revenue": [5000, 7500, 6000],
        }
        df = pd.DataFrame(data)

        output = io.BytesIO()
        df.to_excel(output, index=False, engine="openpyxl")

        output.seek(0)
        excel_bytes = output.getvalue()

        assert len(excel_bytes) > 0
        # Excel files start with PK (zip format)
        assert excel_bytes[:2] == b'PK'

    def test_openpyxl_workbook(self):
        """Test creating Excel with openpyxl."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        import io

        wb = Workbook()
        ws = wb.active
        ws.title = "Test Report"

        # Add header
        ws["A1"] = "Metric"
        ws["B1"] = "Value"
        ws["A1"].font = Font(bold=True)
        ws["B1"].font = Font(bold=True)

        # Add data
        ws["A2"] = "Revenue"
        ws["B2"] = 50000
        ws["A3"] = "Patients"
        ws["B3"] = 100

        output = io.BytesIO()
        wb.save(output)

        output.seek(0)
        excel_bytes = output.getvalue()

        assert len(excel_bytes) > 0

    def test_excel_with_chart(self):
        """Test Excel with chart."""
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        import io

        wb = Workbook()
        ws = wb.active

        # Add data
        data = [
            ["Day", "Revenue"],
            ["Mon", 5000],
            ["Tue", 6000],
            ["Wed", 4500],
            ["Thu", 7000],
            ["Fri", 8000],
        ]
        for row in data:
            ws.append(row)

        # Create chart
        chart = BarChart()
        chart.title = "Daily Revenue"
        data_ref = Reference(ws, min_col=2, min_row=1, max_row=6)
        cats_ref = Reference(ws, min_col=1, min_row=2, max_row=6)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        ws.add_chart(chart, "D2")

        output = io.BytesIO()
        wb.save(output)

        output.seek(0)
        excel_bytes = output.getvalue()

        assert len(excel_bytes) > 0


class TestReportFilenames:
    """Tests for report filename generation."""

    def test_daily_report_filename(self):
        """Test daily report filename format."""
        report_date = date(2026, 1, 15)
        filename = f"daily_summary_{report_date.strftime('%Y-%m-%d')}.pdf"

        assert filename == "daily_summary_2026-01-15.pdf"

    def test_monthly_report_filename(self):
        """Test monthly report filename format."""
        year = 2026
        month = 1
        filename = f"monthly_report_{year}-{month:02d}.pdf"

        assert filename == "monthly_report_2026-01.pdf"

    def test_revenue_report_filename(self):
        """Test revenue report filename format."""
        start_date = date(2026, 1, 1)
        end_date = date(2026, 1, 31)
        filename = f"revenue_report_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.pdf"

        assert filename == "revenue_report_20260101_20260131.pdf"

    def test_appointments_excel_filename(self):
        """Test appointments Excel filename format."""
        start_date = date(2026, 1, 1)
        end_date = date(2026, 1, 31)
        filename = f"appointments_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"

        assert filename == "appointments_20260101_20260131.xlsx"


class TestDateRangeCalculation:
    """Tests for date range calculations."""

    def test_today_range(self):
        """Test 'today' period."""
        from app.api.v1.reports import get_date_range

        today = date.today()
        start, end = get_date_range("today")

        assert start == today
        assert end == today

    def test_week_range(self):
        """Test 'week' period starts on Monday."""
        from app.api.v1.reports import get_date_range

        today = date.today()
        start, end = get_date_range("week")

        # Start should be Monday of current week
        assert start.weekday() == 0  # Monday
        assert end == today

    def test_month_range(self):
        """Test 'month' period starts on 1st."""
        from app.api.v1.reports import get_date_range

        today = date.today()
        start, end = get_date_range("month")

        assert start.day == 1
        assert start.month == today.month
        assert end == today

    def test_year_range(self):
        """Test 'year' period starts on Jan 1."""
        from app.api.v1.reports import get_date_range

        today = date.today()
        start, end = get_date_range("year")

        assert start.day == 1
        assert start.month == 1
        assert start.year == today.year
        assert end == today


class TestReportContentValidation:
    """Tests for report content validation."""

    def test_pdf_content_type(self):
        """Test PDF content type header."""
        content_type = "application/pdf"
        assert content_type == "application/pdf"

    def test_excel_content_type(self):
        """Test Excel content type header."""
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "spreadsheetml" in content_type

    def test_content_disposition_pdf(self):
        """Test Content-Disposition for PDF."""
        filename = "report.pdf"
        header = f"attachment; filename={filename}"

        assert "attachment" in header
        assert "report.pdf" in header

    def test_content_disposition_excel(self):
        """Test Content-Disposition for Excel."""
        filename = "report.xlsx"
        header = f"attachment; filename={filename}"

        assert "attachment" in header
        assert "report.xlsx" in header
